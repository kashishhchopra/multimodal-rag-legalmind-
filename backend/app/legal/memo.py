"""Export an answer plus its citations as a research memo (PDF, DOCX or Markdown)."""
from __future__ import annotations

import io
from typing import Dict, List, Optional

DISCLAIMER = "AI-generated research support. Not legal advice; verify every citation against the source before relying on it."


def _lines(title: str, question: str, answer: str, citations: List[Dict], verification: Optional[Dict]) -> List[tuple]:
    rows = [("title", title), ("h", "Question"), ("p", question), ("h", "Answer"), ("p", answer), ("h", "Sources")]
    for i, c in enumerate(citations, 1):
        rows.append(("p", f"[{i}] {c.get('label', '')} (page {c.get('page_number', '?')}): {c.get('snippet', '')[:300]}"))
    if verification:
        rows += [("h", "Verification"), ("p", f"Status: {verification.get('status')}; claims supported: {verification.get('support_ratio')}")]
        rows += [("p", f"Unsupported claim: {u}") for u in verification.get("unsupported", [])]
    rows += [("h", "Disclaimer"), ("p", DISCLAIMER)]
    return rows


def to_markdown(*args, **kw) -> bytes:
    out = []
    for kind, text in _lines(*args, **kw):
        out.append({"title": f"# {text}", "h": f"\n## {text}"}.get(kind, text))
    return "\n".join(out).encode()


def to_pdf(*args, **kw) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    st = getSampleStyleSheet()
    buf = io.BytesIO()
    story = []
    for kind, text in _lines(*args, **kw):
        style = {"title": st["Title"], "h": st["Heading2"]}.get(kind, st["BodyText"])
        story += [Paragraph(escape(text).replace("\n", "<br/>"), style), Spacer(1, 6)]
    SimpleDocTemplate(buf, pagesize=A4, title="LegalMind memo").build(story)
    return buf.getvalue()


def to_docx(*args, **kw) -> bytes:
    from docx import Document

    doc, buf = Document(), io.BytesIO()
    for kind, text in _lines(*args, **kw):
        if kind == "title":
            doc.add_heading(text, 0)
        elif kind == "h":
            doc.add_heading(text, 2)
        else:
            doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


FORMATS = {"md": (to_markdown, "text/markdown"), "pdf": (to_pdf, "application/pdf"),
           "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
